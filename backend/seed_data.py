import os
from pathlib import Path
import openpyxl
from backend.database import get_db_connection, init_db

BASE_DIR = Path(__file__).resolve().parent.parent
SEED_FILE_PATH = BASE_DIR / "data" / "events_seed.xlsx"

def seed_events():
    """Seeds the events table from data/events_seed.xlsx if the table is empty."""
    # Ensure database tables exist first
    init_db()

    # Step 5: Check if seed file exists
    if not SEED_FILE_PATH.exists():
        raise FileNotFoundError(
            f"Seed data file not found at '{SEED_FILE_PATH}'. "
            "Please copy the provided 'events_seed.xlsx' file into the 'data/' folder before running this script."
        )

    conn = get_db_connection()
    cursor = conn.cursor()

    # Step 4: Skip seeding if the table already has rows
    cursor.execute("SELECT COUNT(*) FROM events")
    count = cursor.fetchone()[0]
    if count > 0:
        print(f"Events table already contains {count} events. Skipping seeding.")
        conn.close()
        return

    # Step 1: Open data/events_seed.xlsx with openpyxl and read the 'Events' sheet
    wb = openpyxl.load_workbook(SEED_FILE_PATH)
    if "Events" not in wb.sheetnames:
        conn.close()
        raise ValueError("Sheet 'Events' not found in events_seed.xlsx.")

    sheet = wb["Events"]
    
    # Step 2: Row 1 is header row. Read rows after it.
    rows = list(sheet.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        conn.close()
        raise ValueError("No data found in 'Events' sheet.")

    header = rows[0]
    expected_header = ['id', 'title', 'category', 'description', 'event_date', 'event_time', 'venue', 'capacity', 'fee', 'icon']
    # Check header columns match
    for idx, col in enumerate(expected_header):
        if idx >= len(header) or str(header[idx]).strip() != col:
            print(f"Warning: Expected header column '{col}' at index {idx}, found '{header[idx] if idx < len(header) else None}'")

    inserted_count = 0
    for row in rows[1:]:
        # Skip completely empty rows if any
        if not any(row):
            continue
        
        event_id = row[0]
        title = row[1]
        category = row[2]
        description = row[3]
        event_date = str(row[4]) if row[4] is not None else ""
        event_time = str(row[5]) if row[5] is not None else ""
        venue = str(row[6]) if row[6] is not None else ""
        capacity = int(row[7]) if row[7] is not None else 0
        fee = int(row[8]) if row[8] is not None else 0
        icon = str(row[9]) if row[9] is not None else ""

        # Step 3: Insert row into events table
        cursor.execute("""
            INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (event_id, title, category, description, event_date, event_time, venue, capacity, fee, icon))
        inserted_count += 1

    conn.commit()
    conn.close()
    print(f"Successfully seeded {inserted_count} events from {SEED_FILE_PATH.name}.")

if __name__ == "__main__":
    seed_events()
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, title, category, venue, capacity, fee, icon FROM events ORDER BY id ASC")
    events = cursor.fetchall()
    print(f"\n--- Events in Database ({len(events)} rows) ---")
    for ev in events:
        safe_icon = ev['icon'].encode('ascii', 'backslashreplace').decode('ascii')
        print(f"[{ev['id']}] {safe_icon} {ev['title']} ({ev['category']}) - Venue: {ev['venue']}, Cap: {ev['capacity']}, Fee: Rs {ev['fee']}")
    conn.close()
